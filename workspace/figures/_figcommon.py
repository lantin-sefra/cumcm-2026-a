"""22 张数据图的共享底座（风格/取色/取数/落盘）。

⛔ 本文件【不是】出图脚本（不以 gen_fig 开头），只被 gen_fig_*.py import。
风格契约：setup_style('nature') 全篇调一次；每张图建好 fig 后立即 set_paper_placement；
取色只走 nature_palette()，禁止 hex 字面量；落盘只走 save_fig（LaTeX 模式 → 仅 PDF）。
"""
from __future__ import annotations

import os
import sys

import numpy as np

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from _utils.plot_utils import (  # noqa: E402
    setup_style, save_fig, set_paper_placement, nature_palette, nature_markers,
    smart_labels, auto_legend, shared_legend, uncertainty_band, dynamic_limits,
    declutter_axes, draw_vector_heatmap,
)

import matplotlib.pyplot as plt  # noqa: E402

setup_style(palette='nature')

# matplotlib 默认往 PDF 里嵌 Type3 字体，中文字形没有 ToUnicode 映射 —— 结果是
# PDF 里的中文完全不可提取：文字体检脚本只看得见数字和拉丁字符，中文标签的
# 重叠/字号一项都没被查过（实测 get_text() 提取到 0 个汉字）。改嵌 TrueType(42)
# 让中文可提取、可搜索、可复制，体检脚本才真正覆盖中文。
import matplotlib  # noqa: E402
matplotlib.rcParams['pdf.fonttype'] = 42
matplotlib.rcParams['ps.fonttype'] = 42

C = nature_palette()
M = nature_markers()
FIGDATA = os.path.join(ROOT, 'figures', '_figdata')
OUTDIR = os.path.join(ROOT, 'figures')

# 交付口径常量（与 code/params.py 一致，仅用于画阈值线/标注，不参与求解）
C_TH = 0.15          # kg/kg 达标阈值
R0_CM = 2.0          # cm 初始半径
SEC_PER_HOUR = 3600.0


def _sheet_array(path, sheet_index=0):
    """将最终 Excel 的一个工作表读成时间、半径与场数组（仅作绘图适配）。"""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[sheet_index]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    radii = np.asarray(header[1:], dtype=float)
    values = [row for row in rows if row[0] is not None]
    t = np.asarray([row[0] for row in values], dtype=float)
    field = np.asarray([row[1:] for row in values], dtype=float)
    wb.close()
    return t, radii, field


def _air_temperature(t):
    """附件1环境温度插值；4 h 后按题目规定平台值延拓。"""
    from openpyxl import load_workbook
    path = os.path.join(ROOT, 'user_data', '附件', '附件', '附件1.xlsx')
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = list(wb.active.iter_rows(min_row=2, values_only=True))
    wb.close()
    tx = np.asarray([row[0] for row in rows], dtype=float)
    temp = np.asarray([row[1] for row in rows], dtype=float)
    return np.where(t > 4.0 * SEC_PER_HOUR, 50.0, np.interp(t, tx, temp))


def _fixed_result(tag):
    """把 result1--result3.xlsx 转为原绘图源码所需的数组键。"""
    number = {'p1': 1, 'p2': 2, 'p3': 3}[tag]
    path = os.path.join(ROOT, 'output', 'result%d.xlsx' % number)
    t, radii, first = _sheet_array(path, 0)
    if number <= 2:
        t2, radii2, second = _sheet_array(path, 1)
        if not np.array_equal(t, t2) or not np.array_equal(radii, radii2):
            raise ValueError('最终 Excel 的两个工作表坐标不一致: %s' % path)
        Tcol, Ccol = first, second
    else:
        Tcol, Ccol = None, first

    data = {
        't': t,
        'cols_cm': radii,
        'Ccol': Ccol,
        'Ccenter': Ccol[:, 0],
        'Csurf': Ccol[:, -1],
        'maxC': np.nanmax(Ccol, axis=1),
        't_end_h': float(t[-1] / SEC_PER_HOUR),
    }
    # 圆柱截面的体平均含水率：2/R^2 * integral(C r dr)。
    data['Cbar'] = 2.0 * np.trapezoid(Ccol * radii[None, :], radii, axis=1) / (radii[-1] ** 2)
    if Tcol is not None:
        data.update({
            'Tcol': Tcol,
            'Tcenter': Tcol[:, 0],
            'Tsurf': Tcol[:, -1],
            'T_air': _air_temperature(t),
        })
    return data


def _moving_result():
    """把 result4.xlsx 与附件2半径记录适配为原移动边界绘图数组。"""
    from openpyxl import load_workbook
    from scipy.interpolate import PchipInterpolator

    path = os.path.join(ROOT, 'output', 'result4.xlsx')
    wb = load_workbook(path, read_only=True, data_only=True)
    result_rows = wb.active.iter_rows(values_only=True)
    header = next(result_rows)
    radii = np.asarray(header[1:-1], dtype=float)
    result_rows = [row for row in result_rows if row[0] is not None]
    t = np.asarray([row[0] for row in result_rows], dtype=float)
    Cfixed = np.asarray([row[1:-1] for row in result_rows], dtype=float)
    Csurface = np.asarray([row[-1] for row in result_rows], dtype=float)
    wb.close()
    radius_path = os.path.join(ROOT, 'user_data', '附件', '附件', '附件2.xlsx')
    wb = load_workbook(radius_path, read_only=True, data_only=True)
    rows = [row for row in wb.active.iter_rows(min_row=2, values_only=True)
            if row[0] is not None and row[1] is not None]
    wb.close()
    rt = np.asarray([row[0] for row in rows], dtype=float)
    rr = np.asarray([row[1] for row in rows], dtype=float)
    R_cm = PchipInterpolator(rt, rr, extrapolate=False)(np.minimum(t, rt[-1]))

    # 原图使用随动坐标；这里只把最终 Excel 的固定半径列插值到同一绘图坐标。
    N = len(radii) - 1
    xi = np.linspace(0.0, 1.0, N + 1)
    nodesC = np.empty((len(t), N + 1), dtype=float)
    for i, radius in enumerate(R_cm):
        inside = radii <= radius + 1e-12
        x = radii[inside]
        y = Cfixed[i, inside]
        finite = np.isfinite(y)
        x, y = x[finite], y[finite]
        surface = Csurface[i]
        if np.isfinite(surface):
            if len(x) == 0 or abs(x[-1] - radius) > 1e-12:
                x = np.append(x, radius)
                y = np.append(y, surface)
            else:
                y[-1] = surface
        nodesC[i] = np.interp(xi * radius, x, y)

    return {
        't': t,
        'R_m': R_cm * 1e-2,
        'nodesC': nodesC,
        'maxC': np.nanmax(nodesC, axis=1),
        'N': N,
        't_end_h': float(t[-1] / SEC_PER_HOUR),
    }


def sha256_file(path):
    """计算数据文件 SHA-256，供出图脚本打印对账。"""
    import hashlib
    digest = hashlib.sha256()
    with open(path, 'rb') as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b''):
            digest.update(chunk)
    return digest.hexdigest()


def describe_source(path, **meta):
    """打印正式数据源路径、SHA-256 与关键口径。"""
    digest = sha256_file(path)
    print('[source] %s' % os.path.relpath(path, ROOT), flush=True)
    print('[source] sha256=%s' % digest, flush=True)
    if meta:
        print('[source] %s' % '  '.join('%s=%s' % item for item in meta.items()),
              flush=True)
    return digest


def load(tag):
    """正式结果图读取最终 Excel；其他历史诊断图仍按原缓存读取。"""
    if tag in ('p1', 'p2', 'p3'):
        number = {'p1': 1, 'p2': 2, 'p3': 3}[tag]
        describe_source(os.path.join(ROOT, 'output', 'result%d.xlsx' % number),
                        tag=tag)
        return _fixed_result(tag)
    if tag == 'p4s2':
        describe_source(os.path.join(ROOT, 'output', 'result4.xlsx'),
                        tag='p4s2', chi=1)
        return _moving_result()
    path = os.path.join(FIGDATA, '%s.npz' % tag)
    if not os.path.isfile(path):
        raise FileNotFoundError(
            '缺少数值场缓存 %s。正式出图应改读 result*.xlsx 或 '
            'output/final_diagnostics/，不要重跑 prep_figdata.py。' % path)
    print('[figdata-legacy] tag=%s 仍读 %s（非正式结果）' % (tag, path),
          flush=True)
    return np.load(path)


def load_json(name):
    """读 output/<name>.json（上游 code/main.py 真算落盘的摘要）。"""
    import json
    path = os.path.join(ROOT, 'output', '%s.json' % name)
    if not os.path.isfile(path):
        raise FileNotFoundError('缺少上游结果 %s' % path)
    with open(path, encoding='utf-8') as f:
        return json.load(f)


def load_csv(name):
    """读 output/<name>（上游真算 CSV）。返回 pandas.DataFrame。"""
    import pandas as pd
    path = os.path.join(ROOT, 'output', name)
    if not os.path.isfile(path):
        raise FileNotFoundError('缺少上游结果 %s' % path)
    return pd.read_csv(path)


def load_formal_csv(name):
    """只读 output/final_diagnostics/ 下的正式诊断，拒绝 legacy。"""
    normalized = name.replace('\\', '/').lstrip('./')
    if normalized.startswith('legacy/') or '/legacy/' in normalized:
        raise RuntimeError('正式出图禁止读取 legacy: %s' % name)
    if not normalized.startswith('final_diagnostics/'):
        raise RuntimeError(
            '正式出图 CSV 必须位于 final_diagnostics/: %s' % name)
    frame = load_csv(normalized)
    describe_source(os.path.join(ROOT, 'output', normalized))
    return frame


def newfig(w, h, width_fraction=0.92):
    """建图并立即登记版面占位（风格契约要求 set_paper_placement 紧随其后）。"""
    fig = plt.figure(figsize=(w, h), layout='constrained')
    # 两处余量都是实测需要的：
    # 1) constrained 按 matplotlib 字体度量排版，PDF 里的字形上升部比它算的再高
    #    1~2 pt，最左 y 标签会压到 mediabox 外；
    # 2) auto_legend 把图例抬进独立行时会重排，没让位的那一行子图会向上扩张，
    #    把 (a)/(b) 面板标签顶出画布 —— 顶部要留够一个面板标签的高度。
    # 左右各留 13 pt：旋转的中文 y 标签横向占位由字高决定，matplotlib 报出的
    # window extent 比 PDF 里的真实字形边缘偏右约 9.75 pt（实测扫墨迹标定：
    # 最左墨迹 ≈ mx − 9.75 pt），所以按它排版必然压边。shared_legend 提交车道后
    # 会 set_layout_engine('none') 冻结轴位、丢掉本 rect，事后无法补救，
    # 必须一开始就留够。
    mx = 13.0 / (w * 72.0)
    my = 2.0 / (h * 72.0)
    top = 6.0 / (h * 72.0)
    fig.get_layout_engine().set(rect=(mx, my, 1.0 - 2.0 * mx, 1.0 - my - top))
    set_paper_placement(fig, width_fraction=width_fraction)
    return fig


def panel(ax, tag):
    """子图面板标签 (a)/(b)/… —— 左上角，非整图标题。"""
    ax.set_title(tag, loc='left', fontweight='bold')


def finish(fig, name):
    """落盘 figures/<name>.pdf（LaTeX 模式仅矢量 PDF，不产 PNG）。"""
    out = os.path.join(OUTDIR, '%s.pdf' % name)
    save_fig(fig, out)
    plt.close(fig)
    print('[fig] %s.pdf' % name, flush=True)


def seq_colors(n, family='blue'):
    """同族 n 级渐变色（时间序/剖面族用）。取色仍源自 nature_palette。"""
    from matplotlib.colors import to_rgb, to_hex
    if family == 'blue':
        a, b = C['blue_secondary'], C['neutral_black']
    elif family == 'red':
        a, b = C['red_1'], C['red_strong']
    elif family == 'green':
        a, b = C['green_1'], C['green_3']
    else:
        a, b = C['neutral_light'], C['neutral_black']
    ra, rb = np.array(to_rgb(a)), np.array(to_rgb(b))
    if n == 1:
        return [to_hex(ra)]
    return [to_hex(ra + (rb - ra) * i / (n - 1.0)) for i in range(n)]


def mono_cmap(family='blue', dark_end=1.0):
    """由 nature_palette 语义色生成单调 colormap（热力图/时空图用）。

    dark_end < 1 时截掉最暗的一段：叠了黑色等值线的时空图里，落进接近纯黑那一端
    的等值线看不见（实测含水率图上端 vmax 处几乎全黑，2.5 那条线整段隐没）。
    """
    import numpy as _np
    from matplotlib.colors import LinearSegmentedColormap
    if family == 'blue':
        stops = ['#FFFFFF', C['blue_secondary'], C['blue_main'], C['neutral_black']]
    elif family == 'red':
        stops = ['#FFFFFF', C['red_1'], C['red_2'], C['red_strong']]
    elif family == 'green':
        stops = ['#FFFFFF', C['green_1'], C['green_2'], C['green_3']]
    else:
        stops = ['#FFFFFF', C['neutral_light'], C['neutral_mid'], C['neutral_black']]
    cm = LinearSegmentedColormap.from_list('mh_%s' % family, stops, N=256)
    if dark_end >= 1.0:
        return cm
    return LinearSegmentedColormap.from_list(
        'mh_%s_t' % family, cm(_np.linspace(0.0, dark_end, 256)), N=256)


def label_levels_on_cbar(cb, cs, fmt='%g'):
    """把等值线层级标到色标上，替代压在热力图上的 inline clabel。

    inline 标注只掏空等值线本身，字仍坐在热力图底色上；落到色标暗端时黑字
    对比度不足（实测 1.x:1），加白底又会撑出 3 倍字高的补丁。改为等值线只留线、
    层级读数移到色标（白底刻度区），既无对比度风险也不遮数据。
    """
    cb.add_lines(cs)
    levels = list(cs.levels)
    cb.set_ticks(levels)
    cb.set_ticklabels([fmt % v for v in levels])
    return cb


def nearest_idx(arr, val):
    """最近邻下标（按报表时刻从密集时间轴取快照，不插值造数）。"""
    return int(np.argmin(np.abs(np.asarray(arr, dtype=float) - float(val))))
